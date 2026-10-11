"""Word and PDF exporters for completed transcripts."""

from __future__ import annotations

from datetime import datetime
from html import escape
from pathlib import Path

from .models import TranscriptionResult
from .utils import format_timestamp
from .views import AUDIO_ONLY, COMBINED, SCREEN_ONLY


def export_reviewer_word(content: str, destination: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10.5)
    document.add_heading("Lecture Reviewer", 0)
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line:
            document.add_paragraph(line)
    document.save(destination)


def export_reviewer_pdf(content: str, destination: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [Paragraph("Lecture Reviewer", styles["Title"]), Spacer(1, 8)]
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if line.startswith("## "):
            story.extend((Spacer(1, 6), Paragraph(escape(line[3:]), styles["Heading1"])))
        elif line.startswith(("- ", "* ")):
            story.append(Paragraph(f"• {escape(line[2:])}", styles["BodyText"]))
        elif line:
            story.append(Paragraph(escape(line), styles["BodyText"]))
    SimpleDocTemplate(
        str(destination), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm, title="Lecture Reviewer"
    ).build(story)


def _view_title(mode: str) -> str:
    return {
        AUDIO_ONLY: "Audio Transcript",
        SCREEN_ONLY: "Screen Context Timeline",
        COMBINED: "Context Aware Transcript",
    }.get(mode, "Context Aware Transcript")


def export_word(
    result: TranscriptionResult, destination: Path, mode: str = COMBINED
) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.color.rgb = None

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Video Lecture Transcript")

    metadata = document.add_paragraph()
    metadata.add_run("Source: ").bold = True
    metadata.add_run(result.video_path.name)
    metadata.add_run("\nLanguage: ").bold = True
    metadata.add_run(result.language)
    metadata.add_run("\nCreated: ").bold = True
    metadata.add_run(datetime.now().strftime("%Y-%m-%d %H:%M"))

    document.add_heading(_view_title(mode), level=1)

    if mode == AUDIO_ONLY:
        segments = result.segments
        if segments:
            for segment in segments:
                paragraph = document.add_paragraph()
                timestamp = paragraph.add_run(
                    f"[{format_timestamp(segment.start)} - {format_timestamp(segment.end)}] "
                )
                timestamp.bold = True
                paragraph.add_run(segment.text)
        else:
            for block in result.blocks:
                paragraph = document.add_paragraph()
                timestamp = paragraph.add_run(
                    f"[{format_timestamp(block.start)} - {format_timestamp(block.end)}] "
                )
                timestamp.bold = True
                paragraph.add_run(block.speech)
        document.save(destination)
        return

    if mode == SCREEN_ONLY:
        if result.screen_contexts:
            entries = [
                (context.timestamp, context.source, context.text)
                for context in result.screen_contexts
            ]
        else:
            entries = [
                (block.start, block.context_source or "Context", block.screen_context)
                for block in result.blocks
                if block.screen_context
            ]
        for timestamp_value, source, text in entries:
            paragraph = document.add_paragraph()
            timestamp = paragraph.add_run(f"[{format_timestamp(timestamp_value)}] ")
            timestamp.bold = True
            label = paragraph.add_run(f"Screen ({source}): ")
            label.bold = True
            paragraph.add_run(text or "")
        document.save(destination)
        return

    for block in result.blocks:
        paragraph = document.add_paragraph()
        timestamp = paragraph.add_run(
            f"[{format_timestamp(block.start)} - {format_timestamp(block.end)}] "
        )
        timestamp.bold = True
        paragraph.add_run(block.speech)

        if block.screen_context:
            context = document.add_paragraph()
            context.paragraph_format.left_indent = Inches(0.25)
            label = context.add_run(f"Screen ({block.context_source}): ")
            label.bold = True
            detail = context.add_run(block.screen_context)
            detail.italic = True

    document.save(destination)


def export_pdf(
    result: TranscriptionResult, destination: Path, mode: str = COMBINED
) -> None:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    document = SimpleDocTemplate(
        str(destination),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="Video Lecture Transcript",
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="Transcript",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            spaceAfter=5,
            alignment=TA_LEFT,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ScreenContext",
            parent=styles["BodyText"],
            fontName="Helvetica-Oblique",
            fontSize=9,
            leading=12,
            leftIndent=12,
            textColor="#3E4C59",
            spaceAfter=8,
        )
    )

    story = [
        Paragraph("Video Lecture Transcript", styles["Title"]),
        Paragraph(f"<b>Source:</b> {escape(result.video_path.name)}", styles["BodyText"]),
        Paragraph(f"<b>Language:</b> {escape(result.language)}", styles["BodyText"]),
        Paragraph(
            f"<b>Created:</b> {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            styles["BodyText"],
        ),
        Spacer(1, 10),
        Paragraph(_view_title(mode), styles["Heading1"]),
    ]

    if mode == AUDIO_ONLY:
        if result.segments:
            entries = [
                (segment.start, segment.end, segment.text) for segment in result.segments
            ]
        else:
            entries = [(block.start, block.end, block.speech) for block in result.blocks]
        for start, end, text in entries:
            time_range = f"{format_timestamp(start)} - {format_timestamp(end)}"
            story.append(
                Paragraph(f"<b>[{time_range}]</b> {escape(text)}", styles["Transcript"])
            )
        document.build(story)
        return

    if mode == SCREEN_ONLY:
        if result.screen_contexts:
            entries = [
                (context.timestamp, context.source, context.text)
                for context in result.screen_contexts
            ]
        else:
            entries = [
                (block.start, block.context_source or "Context", block.screen_context or "")
                for block in result.blocks
                if block.screen_context
            ]
        for timestamp_value, source, text in entries:
            story.append(
                Paragraph(
                    f"<b>[{format_timestamp(timestamp_value)}] Screen ({escape(source)}):</b> "
                    f"{escape(text)}",
                    styles["ScreenContext"],
                )
            )
        document.build(story)
        return

    for block in result.blocks:
        time_range = f"{format_timestamp(block.start)} - {format_timestamp(block.end)}"
        story.append(
            Paragraph(
                f"<b>[{time_range}]</b> {escape(block.speech)}",
                styles["Transcript"],
            )
        )
        if block.screen_context:
            story.append(
                Paragraph(
                    f"<b>Screen ({escape(block.context_source or 'Context')}):</b> "
                    f"{escape(block.screen_context)}",
                    styles["ScreenContext"],
                )
            )

    document.build(story)
